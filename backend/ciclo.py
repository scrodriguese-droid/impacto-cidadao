"""Ciclo econômico e contexto para a política fiscal.

Fontes oficiais (abertas, sem chave):
* IBGE — SIDRA: PIB trimestral (5932), PNAD Contínua (6381, 6390, 6320), Produção Industrial (8888)
  e Pesquisa Mensal de Comércio (8880);
* Banco Central — SGS: IBC-Br com ajuste sazonal (24364), IPCA (433/13522) e séries fiscais;
* Tesouro Nacional — Resultado do Tesouro Nacional (Governo Central, mensal).

Regras: valores preservados como publicados; o que é cálculo nosso vem identificado; a leitura do ciclo
é uma interpretação por critérios explícitos, nunca uma datação oficial de recessão (no Brasil, quem
data os ciclos é o CODACE/FGV IBRE).
"""
from __future__ import annotations

import io
import logging
import re
from datetime import date
from typing import Any

import macro
from fontes import ErroFonte

log = logging.getLogger("impacto.ciclo")

SIDRA = macro.SIDRA


# --------------------------------------------------------------------------- #
# SIDRA
# --------------------------------------------------------------------------- #
def _sidra(tabela: int, variavel: int, filtros: str = "") -> list[dict]:
    """Série completa de uma variável do SIDRA (Brasil). Datas em AAAA-MM-01; trimestres usam o 1º mês."""
    def produzir() -> list[dict]:
        caminho = f"{SIDRA}/t/{tabela}/n1/all/v/{variavel}/{filtros + '/' if filtros else ''}p/all"
        linhas = macro._http.json(caminho, {"formato": "json"})
        pontos: list[dict] = []
        dim = "D4" if filtros else "D3"       # com classificação, o período é a 4ª dimensão
        for r in linhas[1:]:
            cod = str(r.get(f"{dim}C", ""))
            v = macro._num(r.get("V"))
            if v is None or len(cod) != 6:
                continue
            ano, resto = int(cod[:4]), int(cod[4:])
            if tabela == 5932:        # trimestral: 01..04
                if not 1 <= resto <= 4:
                    continue
                mes, rotulo = (resto - 1) * 3 + 1, f"{resto}º tri {ano}"
            else:                     # mensal ou trimestre móvel: AAAAMM
                if not 1 <= resto <= 12:
                    continue
                mes, rotulo = resto, r.get(f"{dim}N")
            pontos.append({"data": f"{ano}-{mes:02d}-01", "valor": v, "referencia": rotulo})
        return sorted(pontos, key=lambda p: p["data"])
    return macro._memo(("sidra", tabela, variavel, filtros), produzir)


def _deriva_variacao(pontos: list[dict], defasagem: int) -> list[dict]:
    """Variação % contra `defasagem` observações antes (cálculo nosso)."""
    saida = []
    for i in range(defasagem, len(pontos)):
        base = pontos[i - defasagem]["valor"]
        if base:
            saida.append({"data": pontos[i]["data"], "valor": round((pontos[i]["valor"] / base - 1) * 100, 2),
                          "referencia": pontos[i].get("referencia")})
    return saida


def _media_3m_var(pontos: list[dict]) -> list[dict]:
    """Variação da média dos 3 últimos meses contra a média dos 3 anteriores (série dessazonalizada)."""
    saida = []
    for i in range(5, len(pontos)):
        a = sum(p["valor"] for p in pontos[i - 2:i + 1]) / 3
        b = sum(p["valor"] for p in pontos[i - 5:i - 2]) / 3
        if b:
            saida.append({"data": pontos[i]["data"], "valor": round((a / b - 1) * 100, 2)})
    return saida


# --------------------------------------------------------------------------- #
# Indicadores
# --------------------------------------------------------------------------- #
def _meta(nome: str, unidade: str, freq: str, fonte: str, url: str, nota: str = "", derivada: bool = False) -> dict:
    return {"nome": nome, "unidade": unidade, "frequencia": freq, "fonte": fonte, "url_fonte": url, "nota": nota, "calculo_proprio": derivada}


def indicadores() -> dict[str, dict]:
    sidra = lambda t: f"{SIDRA}/t/{t}"  # noqa: E731
    pib_qoq = _sidra(5932, 6564, "c11255/90707")
    pib_yoy = _sidra(5932, 6561, "c11255/90707")
    pib_4t = _sidra(5932, 6562, "c11255/90707")
    ibc = macro._bcb(24364, 2012)
    pim = _sidra(8888, 12607, "c544/129314")           # número-índice com ajuste sazonal (2022=100)
    pmc = _sidra(8880, 7170, "c11046/56734")           # volume de vendas, com ajuste sazonal
    des = macro._sidra_desemprego()
    renda = _sidra(6390, 5933)
    ocup = _sidra(6320, 4090, "c11913/96165")
    ipca12 = macro._bcb(13522, 2012)

    ind: dict[str, dict] = {}
    ind["pib_qoq"] = {**_meta("PIB — variação contra o trimestre anterior (com ajuste sazonal)", "% sobre o trimestre anterior", "trimestral",
                              "IBGE — Contas Nacionais Trimestrais (SIDRA 5932, variável 6564)", sidra(5932)), "pontos": pib_qoq}
    ind["pib_yoy"] = {**_meta("PIB — variação contra o mesmo trimestre do ano anterior", "% sobre o mesmo trimestre do ano anterior", "trimestral",
                              "IBGE — Contas Nacionais Trimestrais (SIDRA 5932, variável 6561)", sidra(5932)), "pontos": pib_yoy}
    ind["pib_4t"] = {**_meta("PIB — acumulado em quatro trimestres", "% (4 trimestres contra os 4 anteriores)", "trimestral",
                             "IBGE — Contas Nacionais Trimestrais (SIDRA 5932, variável 6562)", sidra(5932)), "pontos": pib_4t}
    ind["ibc_br"] = {**_meta("IBC-Br — índice com ajuste sazonal", "índice (2002=100)", "mensal",
                             "Banco Central — SGS 24364", f"{macro.SGS}.24364/dados?formato=json",
                             "Proxy mensal do PIB; é uma estimativa do Banco Central, não o PIB oficial do IBGE."), "pontos": ibc}
    ind["ibc_br_mm"] = {**_meta("IBC-Br — variação mensal (com ajuste sazonal)", "% sobre o mês anterior", "mensal",
                                "Calculado a partir do SGS 24364", f"{macro.SGS}.24364/dados?formato=json", "Cálculo nosso.", True),
                        "pontos": _deriva_variacao(ibc, 1)}
    ind["desocupacao"] = {**_meta("Taxa de desocupação (PNAD Contínua)", "% da força de trabalho", "trimestre móvel, mensal",
                                  "IBGE — PNAD Contínua (SIDRA 6381, variável 4099)", sidra(6381)), "pontos": des}
    ind["rendimento"] = {**_meta("Rendimento médio real habitual (PNAD Contínua)", "R$ por mês (preços do último trimestre)", "trimestre móvel, mensal",
                                 "IBGE — PNAD Contínua (SIDRA 6390, variável 5933)", sidra(6390)), "pontos": renda}
    ind["ocupados"] = {**_meta("Pessoas ocupadas (PNAD Contínua)", "mil pessoas", "trimestre móvel, mensal",
                               "IBGE — PNAD Contínua (SIDRA 6320, variável 4090)", sidra(6320)), "pontos": ocup}
    ind["producao_industrial"] = {**_meta("Produção industrial — indústria geral, com ajuste sazonal", "número-índice (2022=100)", "mensal",
                                          "IBGE — Pesquisa Industrial Mensal (SIDRA 8888, variável 12607)", sidra(8888)), "pontos": pim}
    ind["producao_industrial_mm"] = {**_meta("Produção industrial — variação mensal (com ajuste sazonal)", "% sobre o mês anterior", "mensal",
                                             "IBGE — Pesquisa Industrial Mensal (SIDRA 8888, variável 11601)", sidra(8888)),
                                     "pontos": _sidra(8888, 11601, "c544/129314")}
    ind["vendas_varejo"] = {**_meta("Vendas do comércio varejista — volume, com ajuste sazonal", "número-índice (2022=100)", "mensal",
                                    "IBGE — Pesquisa Mensal de Comércio (SIDRA 8880, variável 7170)", sidra(8880)), "pontos": pmc}
    ind["vendas_varejo_mm"] = {**_meta("Vendas do comércio varejista — variação mensal (com ajuste sazonal)", "% sobre o mês anterior", "mensal",
                                       "IBGE — Pesquisa Mensal de Comércio (SIDRA 8880, variável 11708)", sidra(8880)),
                               "pontos": _sidra(8880, 11708, "c11046/56734")}
    ind["ipca_12m"] = {**_meta("IPCA — acumulado em 12 meses", "% em 12 meses", "mensal",
                               "Banco Central — SGS 13522 (IBGE)", f"{macro.SGS}.13522/dados?formato=json"), "pontos": ipca12}
    return ind


# --------------------------------------------------------------------------- #
# Leitura do ciclo (interpretação por critérios explícitos)
# --------------------------------------------------------------------------- #
LIMIAR_PIB = 0.1        # % por trimestre
LIMIAR_ATIV = 0.2       # % na média móvel de 3 meses
LIMIAR_EMP = 0.3        # % de ocupados em 12 meses
LIMIAR_DESOC_PP = 0.5   # pontos percentuais em 12 meses
LIMIAR_QUEDA_FORTE_ATIV = -1.0   # variação do índice em 6 meses
LIMIAR_QUEDA_FORTE_PIB = -0.5    # soma das taxas dos 2 últimos trimestres


def _duracao_queda(valores: list[float], limiar: float) -> int:
    n = 0
    for v in reversed(valores):
        if v < -limiar:
            n += 1
        else:
            break
    return n


def _classe(v: float, limiar: float) -> str:
    return "avanco" if v > limiar else "queda" if v < -limiar else "estavel"


def leitura_ciclo(ind: dict[str, dict]) -> dict:
    sinais: list[dict] = []

    qoq = ind["pib_qoq"]["pontos"]
    yoy = ind["pib_yoy"]["pontos"]
    if qoq:
        ult = qoq[-1]
        vals = [p["valor"] for p in qoq[-8:]]
        sinais.append({
            "id": "pib", "nome": "PIB trimestral", "classe": _classe(ult["valor"], LIMIAR_PIB),
            "valor": ult["valor"], "unidade": "% sobre o trimestre anterior (c/ ajuste sazonal)", "referencia": ult.get("referencia"),
            "duracao_quedas": _duracao_queda(vals, LIMIAR_PIB),
            "intensidade": round(sum(p["valor"] for p in qoq[-2:]), 2), "intensidade_rotulo": "acumulado nos 2 últimos trimestres (soma simples das taxas)",
        })

    for chave, nome, serie_idx in (("ibc_br", "IBC-Br (proxy mensal do PIB)", "ibc_br"),
                                   ("producao_industrial", "Produção industrial", "producao_industrial"),
                                   ("vendas_varejo", "Vendas do comércio varejista", "vendas_varejo")):
        pts = ind[serie_idx]["pontos"]
        var = _media_3m_var(pts)
        if not var:
            continue
        ult = var[-1]
        sinais.append({
            "id": chave, "nome": nome, "classe": _classe(ult["valor"], LIMIAR_ATIV), "valor": ult["valor"],
            "unidade": "% — média dos 3 últimos meses contra os 3 anteriores (c/ ajuste sazonal)", "referencia": ult["data"][:7],
            "duracao_quedas": _duracao_queda([v["valor"] for v in var[-12:]], LIMIAR_ATIV),
            "intensidade": round((pts[-1]["valor"] / pts[-7]["valor"] - 1) * 100, 2) if len(pts) >= 7 else None,
            "intensidade_rotulo": "variação do índice em 6 meses",
        })

    ocup = ind["ocupados"]["pontos"]
    des = ind["desocupacao"]["pontos"]
    if len(ocup) > 12 and len(des) > 12:
        var_oc = round((ocup[-1]["valor"] / ocup[-13]["valor"] - 1) * 100, 2)
        var_de = round(des[-1]["valor"] - des[-13]["valor"], 2)
        if var_oc < -LIMIAR_EMP or var_de > LIMIAR_DESOC_PP:
            classe = "queda"
        elif var_oc > LIMIAR_EMP and var_de <= LIMIAR_DESOC_PP:
            classe = "avanco"
        else:
            classe = "estavel"
        sinais.append({
            "id": "emprego", "nome": "Mercado de trabalho", "classe": classe, "valor": var_oc,
            "unidade": "% de pessoas ocupadas em 12 meses", "referencia": ocup[-1].get("referencia"),
            "duracao_quedas": 0, "intensidade": var_de, "intensidade_rotulo": "variação da taxa de desocupação em 12 meses (p.p.)",
        })

    for s_ in sinais:
        lim = LIMIAR_QUEDA_FORTE_PIB if s_["id"] == "pib" else LIMIAR_QUEDA_FORTE_ATIV
        s_["queda_intensa"] = bool(s_["classe"] == "queda" and s_["id"] != "emprego"
                                   and s_.get("intensidade") is not None and s_["intensidade"] <= lim)
    avancos = sum(1 for s in sinais if s["classe"] == "avanco")
    quedas = sum(1 for s in sinais if s["classe"] == "queda")
    estaveis = len(sinais) - avancos - quedas
    persistentes = sum(1 for s in sinais if s["duracao_quedas"] >= 2)
    intensas = sum(1 for s in sinais if s["queda_intensa"])
    pib_q = next((s for s in sinais if s["id"] == "pib"), None)
    desacelerando = len(yoy) >= 5 and 0 < yoy[-1]["valor"] < yoy[-5]["valor"] - 0.3

    criterios: list[str] = []
    if quedas >= 3 and persistentes >= 2 and intensas >= 2:
        situacao = "possivel_contracao"
        criterios.append(f"{quedas} de {len(sinais)} sinais em queda (disseminação).")
        criterios.append(f"{persistentes} indicadores com queda por 2 períodos seguidos ou mais (duração).")
        criterios.append(f"{intensas} indicadores com queda acumulada intensa (intensidade).")
    elif avancos >= 3 and quedas == 0 and (pib_q is None or pib_q["valor"] > 0):
        situacao = "expansao"
        criterios.append(f"{avancos} de {len(sinais)} sinais em avanço e nenhum em queda.")
        if pib_q:
            criterios.append("PIB cresceu contra o trimestre anterior.")
    elif quedas >= 1 or estaveis >= 2 or desacelerando:
        situacao = "desaceleracao"
        if quedas:
            faltando = []
            if quedas < 3:
                faltando.append("disseminação")
            if persistentes < 2:
                faltando.append("duração")
            if intensas < 2:
                faltando.append("intensidade")
            criterios.append(f"{quedas} sinal(is) em queda, mas sem {', '.join(faltando) or 'os requisitos'} de uma contração "
                             f"(hoje: {quedas} em queda, {persistentes} persistentes, {intensas} intensos).")
        if estaveis >= 2:
            criterios.append(f"{estaveis} sinais estáveis (variação dentro da margem).")
        if desacelerando:
            criterios.append("O crescimento do PIB contra o mesmo trimestre do ano anterior ficou menor que um ano antes.")
    else:
        situacao = "indeterminado"
        criterios.append("Os sinais disponíveis não se encaixam claramente em nenhuma das três situações.")

    return {
        "situacao": situacao,
        "criterios_atendidos": criterios,
        "contagem": {"avanco": avancos, "estavel": estaveis, "queda": quedas, "total": len(sinais), "queda_persistente": persistentes},
        "sinais": sinais,
        "regras": {
            "limiares": {
                "pib_trimestral_pct": LIMIAR_PIB, "atividade_mensal_media_3m_pct": LIMIAR_ATIV,
                "ocupados_12m_pct": LIMIAR_EMP, "desocupacao_12m_pp": LIMIAR_DESOC_PP,
                "queda_intensa_atividade_6m_pct": LIMIAR_QUEDA_FORTE_ATIV, "queda_intensa_pib_2tri_pct": LIMIAR_QUEDA_FORTE_PIB,
            },
            "descricao": {
                "expansao": "Pelo menos 3 sinais em avanço, nenhum em queda e PIB positivo contra o trimestre anterior.",
                "desaceleracao": "Sinais mistos: ao menos um em queda, ou ao menos dois estáveis, ou PIB em ritmo menor que um ano antes.",
                "possivel_contracao": "Reúne as três condições: ao menos 3 sinais em queda (disseminação), ao menos 2 com queda por 2 ou mais períodos seguidos (duração) "
                                      "e ao menos 2 com queda acumulada intensa (intensidade).",
            },
            "aviso": "Leitura interpretativa por regras simples e abertas. Não é a datação oficial de recessão, que no Brasil cabe ao CODACE/FGV IBRE, "
                     "e nenhum indicador isolado a define. Os limiares são convenções nossas.",
        },
    }


# --------------------------------------------------------------------------- #
# Contas públicas: Resultado do Tesouro Nacional (Governo Central)
# --------------------------------------------------------------------------- #
def _rtn() -> dict:
    def produzir() -> dict:
        from openpyxl import load_workbook

        rec = macro._recurso("resultado-do-tesouro-nacional", "XLSX")
        bruto = b""
        for _ in range(4):
            bruto = macro._http_arquivos.get(rec["url"], accept="*/*").content
            if bruto[:2] == b"PK":
                break
        if bruto[:2] != b"PK":
            raise ErroFonte("Arquivo do Resultado do Tesouro Nacional indisponível")
        ws = load_workbook(io.BytesIO(bruto), read_only=True, data_only=True)["1.2"]
        linhas = list(ws.iter_rows(values_only=True))
        cab = next(r for r in linhas if r and r[0] == "Discriminação")
        colunas = [(i, c.strftime("%Y-%m")) for i, c in enumerate(cab) if hasattr(c, "strftime")]

        def serie(prefixo: str) -> dict[str, float]:
            alvo = next((r for r in linhas if isinstance(r[0], str) and re.sub(r"\s+", " ", r[0]).strip().startswith(prefixo)), None)
            if alvo is None:
                return {}
            return {m: float(alvo[i]) for i, m in colunas if isinstance(alvo[i], (int, float))}

        codigos: dict[str, dict] = {}
        for r in linhas:
            if not isinstance(r[0], str):
                continue
            m = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.*)$", re.sub(r"\s+", " ", r[0]).strip())
            if m and m.group(1) not in codigos:
                codigos[m.group(1)] = {"rotulo": m.group(2), "serie": {mes: float(r[i]) for i, mes in colunas if isinstance(r[i], (int, float))}}
        return {
            "codigos": codigos,
            "receita_total": serie("1. RECEITA TOTAL"),
            "receita_rfb": serie("1.1 - Receita Administrada pela RFB"),
            "receita_liquida": serie("3. RECEITA LÍQUIDA"),
            "despesa_total": serie("4. DESPESA TOTAL"),
            "abono_seguro": serie("4.3.01 Abono e Seguro Desemprego"),
            "resultado_primario": serie("7. RESULTADO PRIMÁRIO") or serie("6. RESULTADO PRIMÁRIO") or serie("5. RESULTADO PRIMÁRIO"),
            "atualizado": rec["atualizado"], "pagina": rec["pagina"], "url": rec["url"],
        }
    return macro._memo(("rtn",), produzir)


def _soma_12(serie: dict[str, float], fim: str, deslocamento: int = 0) -> float | None:
    ano, mes = int(fim[:4]), int(fim[5:7])
    idx = ano * 12 + mes - 1 - deslocamento
    valores = []
    for k in range(12):
        j = idx - k
        chave = f"{j // 12}-{j % 12 + 1:02d}"
        if chave not in serie:
            return None
        valores.append(serie[chave])
    return round(sum(valores), 2)


def contas_publicas(ind: dict[str, dict]) -> dict:
    out: dict[str, Any] = {"disponivel": False}
    try:
        rtn = _rtn()
    except Exception as exc:  # noqa: BLE001
        log.warning("RTN indisponível: %s", exc)
        out["erro"] = "Resultado do Tesouro Nacional indisponível no momento."
        return out
    meses = sorted(rtn["receita_total"])
    if not meses:
        return out
    fim = meses[-1]
    ipca12 = {p["data"][:7]: p["valor"] for p in ind["ipca_12m"]["pontos"]}
    itens = []
    for chave, nome in (("receita_total", "Receita total (arrecadação do Governo Central)"),
                        ("receita_rfb", "Receita administrada pela Receita Federal"),
                        ("despesa_total", "Despesa total (primária)"),
                        ("resultado_primario", "Resultado primário do Governo Central")):
        s = rtn.get(chave) or {}
        a, b = _soma_12(s, fim), _soma_12(s, fim, 12)
        nominal = round((a / b - 1) * 100, 1) if a is not None and b else None
        real = None
        if nominal is not None and ipca12.get(fim) is not None and chave != "resultado_primario":
            real = round(((1 + nominal / 100) / (1 + ipca12[fim] / 100) - 1) * 100, 1)
        itens.append({"id": chave, "nome": nome, "ate": fim, "ultimos_12m": a, "12m_anteriores": b,
                      "variacao_nominal_pct": nominal if chave != "resultado_primario" else None, "variacao_real_aprox_pct": real})
    out.update({
        "disponivel": True, "periodo": fim, "itens": itens, "unidade": "R$ milhões, valores correntes",
        "serie_mensal": [{"mes": m, "receita_rfb": rtn["receita_rfb"].get(m), "abono_seguro": rtn["abono_seguro"].get(m),
                          "receita_total": rtn["receita_total"].get(m), "despesa_total": rtn["despesa_total"].get(m)}
                         for m in meses[-72:]],
        "fonte": {"nome": "Tesouro Nacional — Resultado do Tesouro Nacional (série histórica mensal)", "url": rtn["pagina"], "atualizado": rtn["atualizado"]},
        "nota": "Governo Central (Tesouro, Previdência e Banco Central), não o setor público consolidado. A variação real é cálculo nosso: variação nominal deflacionada pelo IPCA de 12 meses.",
    })
    return out


def relacao_fiscal() -> dict:
    """PIB (trimestral) × resultado primário e nominal em 12 meses (% do PIB) e dívida, séries do Banco Central."""
    chaves = {"primario12": "Resultado primário, 12 meses (% do PIB)", "nominal12": "Resultado nominal, 12 meses (% do PIB)",
              "dbgg_pib": "Dívida bruta do governo geral (% do PIB)", "dlsp_pib": "Dívida líquida do setor público (% do PIB)"}
    saida: dict[str, Any] = {"series": {}}
    for k, nome in chaves.items():
        try:
            pts = macro.serie(k)["pontos"]
        except Exception as exc:  # noqa: BLE001
            log.warning("série %s indisponível: %s", k, exc)
            continue
        saida["series"][k] = {"nome": nome, "pontos": [{"data": p["data"], "valor": p["valor"]} for p in pts[-96:]]}
    return saida


def ciclo() -> dict:
    ind = indicadores()
    leitura = leitura_ciclo(ind)
    cartoes = []
    for chave in ("pib_qoq", "pib_yoy", "ibc_br_mm", "desocupacao", "producao_industrial_mm", "vendas_varejo_mm", "ipca_12m"):
        i = ind[chave]
        if i["pontos"]:
            u = i["pontos"][-1]
            cartoes.append({"id": chave, "nome": i["nome"], "unidade": i["unidade"], "valor": u["valor"], "data": u["data"],
                            "referencia": u.get("referencia"),
                            "anterior": i["pontos"][-2]["valor"] if len(i["pontos"]) > 1 else None,
                            "frequencia": i["frequencia"], "fonte": i["fonte"], "url_fonte": i["url_fonte"]})
    return {
        "consultado_em": macro._agora(),
        "cartoes": cartoes,
        "leitura": leitura,
        "indicadores": {k: {kk: vv for kk, vv in v.items() if kk != "pontos"} | {"pontos": v["pontos"][-180:]} for k, v in ind.items()},
        "contas_publicas": contas_publicas(ind),
        "relacao_fiscal": relacao_fiscal(),
        "revisoes": ("O PIB trimestral é revisado a cada divulgação, e as séries dessazonalizadas (IBC-Br, produção industrial, vendas) "
                     "são reestimadas e podem mudar valores passados. A PNAD usa trimestres móveis: cada ponto cobre três meses."),
    }
