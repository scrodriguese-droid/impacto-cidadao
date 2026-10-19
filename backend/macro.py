"""Indicadores macroeconômicos e fiscais do Brasil — séries oficiais, sem estimativas nossas.

Fontes (todas abertas, sem chave):
* Banco Central do Brasil — SGS  https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados
* IBGE — SIDRA                   https://apisidra.ibge.gov.br  (PNAD Contínua, tabela 6381)
* Tesouro Nacional — Tesouro Transparente (CKAN)  https://www.tesourotransparente.gov.br

Regras de qualidade: valores e unidades preservados; nada é interpolado ou estimado; cada série
leva a URL da fonte, a data da consulta e o período de referência. Os códigos do SGS foram
conferidos nos metadados do BCB (dadosabertos.bcb.gov.br) e entre si: nominal = primário + juros.
"""
from __future__ import annotations

import csv
import io
import logging
import re
import threading
import time
from collections import defaultdict
from datetime import date, datetime, timezone
from typing import Any

from fontes import DISJUNTOR_BCB, ClienteHTTP, ErroFonte

log = logging.getLogger("impacto.macro")

_http = ClienteHTTP(pausa=0.05, timeout=60.0)
_http_sgs = ClienteHTTP(pausa=0.05, timeout=12.0, tentativas=1)   # o SGS segura pedidos recusados por ~20 s: não esperamos tanto
_http_arquivos = ClienteHTTP(pausa=0.05, timeout=240.0)  # CSV/XLSX do Tesouro (10–15 MB)

TTL_S = 6 * 3600
_cache: dict[tuple, tuple[float, Any]] = {}

SGS = "https://api.bcb.gov.br/dados/serie/bcdata.sgs"
CKAN = "https://www.tesourotransparente.gov.br/ckan/api/3/action"
SIDRA = "https://apisidra.ibge.gov.br/values"
PAGINA_SGS = "https://www3.bcb.gov.br/sgspub/"
PAGINA_TESOURO = "https://www.tesourotransparente.gov.br/ckan/dataset/"


def _agora() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


_travas: dict[tuple, threading.Lock] = defaultdict(threading.Lock)


def _memo(chave: tuple, produzir):  # noqa: ANN001
    """Cache com TTL; requisições simultâneas à mesma fonte esperam uma só consulta (sem baixar 2x)."""
    agora = time.time()
    if chave in _cache and agora - _cache[chave][0] < TTL_S:
        return _cache[chave][1]
    with _travas[chave]:
        agora = time.time()
        if chave in _cache and agora - _cache[chave][0] < TTL_S:
            return _cache[chave][1]
        valor = produzir()
        _cache[chave] = (agora, valor)
        return valor


# --------------------------------------------------------------------------- #
# Catálogo de séries
# --------------------------------------------------------------------------- #
FISCAL_NOTA = ("Setor público consolidado (Governo Central + estados + municípios + empresas estatais, "
               "exceto Petrobras e Eletrobras), conceito NFSP “sem desvalorização cambial”. "
               "Sinal do BCB: positivo = déficit (necessidade de financiamento); negativo = superávit.")

NOTA_FLUXO_BPM = ('Balanço de Pagamentos: são as TRANSAÇÕES do mês (variação por compras, vendas e rendimentos), não o saldo. O BCB não publica o saldo desta rubrica por série aberta.')

CATALOGO: dict[str, dict] = {
    "ipca": dict(fonte="bcb", codigo=433, grupo="Inflação", nome="IPCA — variação mensal", unidade="% ao mês",
                 frequencia="mensal", tipo="taxa",
                 mede="Variação mensal do índice oficial de preços ao consumidor (IBGE).",
                 interpreta="Mostra a inflação do mês. Compare com a meta de inflação e com o acumulado em 12 meses."),
    "ipca12": dict(fonte="bcb", codigo=13522, grupo="Inflação", nome="IPCA — acumulado em 12 meses", unidade="% em 12 meses",
                   frequencia="mensal", tipo="taxa",
                   mede="Inflação acumulada nos últimos 12 meses (a referência da meta de inflação).",
                   interpreta="Suaviza a sazonalidade. Difere do IPCA mensal, que mostra só o mês."),
    "selic": dict(fonte="bcb", codigo=432, grupo="Juros", nome="Meta da taxa Selic", unidade="% ao ano",
                  frequencia="diária (muda a cada reunião do Copom)", tipo="taxa", alteracoes=True,
                  mede="Meta para a taxa básica de juros definida pelo Copom.",
                  interpreta="Juros mais altos encarecem o crédito e tendem a esfriar a economia e a inflação; também elevam o custo da dívida pública."),
    "desemprego": dict(fonte="sidra", codigo=6381, grupo="Mercado de trabalho", nome="Taxa de desocupação (PNAD Contínua)",
                       unidade="% da força de trabalho", frequencia="trimestre móvel, divulgada todo mês", tipo="taxa",
                       mede="Pessoas de 14 anos ou mais que procuraram trabalho e não encontraram, como % da força de trabalho.",
                       interpreta="Cada observação cobre um trimestre móvel (ex.: jul-ago-set). Não é o mesmo que o desemprego mensal de outras fontes."),
    "ibcbr": dict(fonte="bcb", codigo=24364, grupo="Atividade econômica", nome="IBC-Br — índice de atividade (dessazonalizado)",
                  unidade="índice (2002 = 100)", frequencia="mensal", tipo="taxa",
                  mede="Prévia mensal do nível de atividade da economia, calculada pelo Banco Central, já sem o efeito das estações do ano.",
                  interpreta="Serve de termômetro do PIB entre uma divulgação trimestral e outra. Não é o PIB oficial do IBGE."),
    "primario": dict(fonte="bcb", codigo=4649, grupo="Resultado fiscal", nome="Resultado primário — fluxo mensal",
                     unidade="R$ milhões", frequencia="mensal", tipo="valor", nota=FISCAL_NOTA,
                     mede="Receitas menos despesas, sem contar juros. Mostra se o governo gasta mais do que arrecada antes de pagar a dívida.",
                     interpreta="Déficit primário = a dívida cresce mesmo sem contar juros. Difere do nominal, que inclui os juros."),
    "juros": dict(fonte="bcb", codigo=4616, grupo="Resultado fiscal", nome="Juros nominais — fluxo mensal",
                  unidade="R$ milhões", frequencia="mensal", tipo="valor", nota=FISCAL_NOTA,
                  mede="Juros pagos e incorridos sobre a dívida pública, incluindo a correção monetária.",
                  interpreta="Sozinho não é déficit: somado ao resultado primário dá o resultado nominal."),
    "nominal": dict(fonte="bcb", codigo=4583, grupo="Resultado fiscal", nome="Resultado nominal — fluxo mensal",
                    unidade="R$ milhões", frequencia="mensal", tipo="valor", nota=FISCAL_NOTA,
                    mede="Resultado primário + juros nominais: quanto o setor público precisou se financiar no mês.",
                    interpreta="É o que faz o estoque da dívida crescer (além de efeitos de câmbio e outros ajustes)."),
    "primario12": dict(fonte="bcb", codigo=5793, grupo="Resultado fiscal", nome="Resultado primário — 12 meses, % do PIB",
                       unidade="% do PIB (acumulado 12 meses)", frequencia="mensal", tipo="taxa", nota=FISCAL_NOTA,
                       mede="Soma dos últimos 12 meses do resultado primário, em proporção do PIB.",
                       interpreta="Permite comparar anos e países, ao contrário do valor em reais."),
    "nominal12": dict(fonte="bcb", codigo=5727, grupo="Resultado fiscal", nome="Resultado nominal — 12 meses, % do PIB",
                      unidade="% do PIB (acumulado 12 meses)", frequencia="mensal", tipo="taxa", nota=FISCAL_NOTA,
                      mede="Soma dos últimos 12 meses do resultado nominal, em proporção do PIB.",
                      interpreta="Inclui os juros: costuma ser bem pior que o primário em anos de juros altos."),
    "juros12": dict(fonte="bcb", codigo=5760, grupo="Resultado fiscal", nome="Juros nominais — 12 meses, % do PIB",
                    unidade="% do PIB (acumulado 12 meses)", frequencia="mensal", tipo="taxa", nota=FISCAL_NOTA,
                    mede="Soma dos últimos 12 meses dos juros nominais, em proporção do PIB.",
                    interpreta="Mostra o peso dos juros na economia."),
    "dlsp_pib": dict(fonte="bcb", codigo=4513, grupo="Dívida pública", nome="Dívida líquida do setor público — % do PIB",
                     unidade="% do PIB (saldo)", frequencia="mensal", tipo="taxa",
                     nota="Dívida LÍQUIDA do setor público consolidado: desconta ativos (reservas, créditos). Não é o mesmo que a Dívida Pública Federal (bruta, só da União).",
                     mede="Saldo da dívida do setor público descontados os ativos financeiros, como % do PIB.",
                     interpreta="Estoque, não fluxo: sobe com déficits nominais. Não confundir com a dívida bruta do governo geral nem com a DPF."),
    "dlsp_rs": dict(fonte="bcb", codigo=4478, grupo="Dívida pública", nome="Dívida líquida do setor público — R$ milhões",
                    unidade="R$ milhões (saldo)", frequencia="mensal", tipo="valor",
                    nota="Dívida LÍQUIDA do setor público consolidado.",
                    mede="Saldo em reais da dívida líquida do setor público.", interpreta="Mesmo indicador do % do PIB, em reais correntes."),
    "dbgg_pib": dict(fonte="bcb", codigo=13762, grupo="Dívida pública", nome="Dívida bruta do governo geral — % do PIB",
                     unidade="% do PIB (saldo)", frequencia="mensal", tipo="taxa",
                     nota="Dívida BRUTA do governo geral (União, estados e municípios; metodologia a partir de 2008). Não desconta ativos; difere da dívida líquida do setor público.",
                     mede="Saldo bruto da dívida do governo geral como proporção do PIB.",
                     interpreta="É a medida de dívida mais usada em comparações internacionais."),
    # Reservas internacionais (estoque, US$ milhões) e transações do Balanço de Pagamentos
    "reservas_caixa": dict(fonte="bcb", codigo=3546, grupo="Reservas internacionais",
                           nome="Reservas internacionais — conceito caixa (fim do mês)", unidade="US$ milhões",
                           frequencia="mensal", tipo="valor",
                           nota="Saldo (estoque), em dólares. Coincide com o conceito caixa diário (SGS 13621) no fim de cada mês.",
                           mede="Ativos externos de liquidez imediata sob controle do Banco Central: títulos, depósitos, ouro, DES e posição no FMI.",
                           interpreta="Funcionam como seguro contra crises cambiais. Variam com compras e vendas de dólares, rendimentos e a cotação dos ativos."),
    "reservas_liquidez": dict(fonte="bcb", codigo=13982, grupo="Reservas internacionais",
                              nome="Reservas internacionais — conceito liquidez (diária)", unidade="US$ milhões",
                              frequencia="diária", tipo="valor",
                              nota="Conceito liquidez: o que o BCB divulga como “reservas internacionais” no dia a dia; é maior que o conceito caixa.",
                              mede="Reservas internacionais pelo conceito liquidez internacional, saldo diário.",
                              interpreta="Mesma ideia do conceito caixa, com abrangência maior. Não compare valores dos dois conceitos entre si."),
    "fmi_posicao_fluxo": dict(fonte="bcb", codigo=23048, grupo="Reservas internacionais",
                              nome="Posição de reserva no FMI — transações do mês", unidade="US$ milhões (fluxo)",
                              frequencia="mensal", tipo="valor", nota=NOTA_FLUXO_BPM,
                              mede="Variação mensal do crédito do Brasil junto ao FMI contado como reserva.",
                              interpreta="Positivo/negativo indica aumento/redução da posição no mês; não informa o saldo."),
    "des_fluxo": dict(fonte="bcb", codigo=23047, grupo="Reservas internacionais",
                      nome="Direitos Especiais de Saque (DES) — transações do mês", unidade="US$ milhões (fluxo)",
                      frequencia="mensal", tipo="valor", nota=NOTA_FLUXO_BPM,
                      mede="Variação mensal da posse de DES, o ativo de reserva criado pelo FMI.",
                      interpreta="Mostra movimentos, não o saldo de DES."),
    "ouro_fluxo": dict(fonte="bcb", codigo=23044, grupo="Reservas internacionais",
                       nome="Ouro monetário — transações do mês", unidade="US$ milhões (fluxo)",
                       frequencia="mensal", tipo="valor", nota=NOTA_FLUXO_BPM,
                       mede="Variação mensal do ouro mantido como reserva.", interpreta="Zero indica que não houve compra ou venda no mês."),
}
# Séries mensais podem ir bem longe; a Selic (diária) é buscada em janelas de até 10 anos.
INICIO_PADRAO = date(2010, 1, 1)


def _iso(d: str) -> str:
    dia, mes, ano = d.split("/")
    return f"{ano}-{mes}-{dia}"


def _num(txt: Any) -> float | None:
    try:
        return float(str(txt).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _json_com_nova_tentativa(url: str, params: dict) -> Any:
    """O SGS às vezes devolve uma página de erro com status 200; tenta mais uma vez antes de desistir.
    Se as duas falharem, o disjuntor abre e as próximas chamadas falham na hora (ver `Disjuntor`)."""
    DISJUNTOR_BCB.exigir_fechado()
    ultimo: Exception | None = None
    for tentativa in range(2):
        if tentativa:
            time.sleep(1)
        try:
            dados = _http_sgs.json(url, params)
            DISJUNTOR_BCB.ok()
            return dados
        except ErroFonte as exc:
            ultimo = exc
    DISJUNTOR_BCB.falhou()
    raise ErroFonte(str(ultimo))


def _bcb(codigo: int, desde: int | None = None) -> list[dict]:
    """Série do SGS desde `desde` (padrão INICIO_PADRAO), em janelas ≤ 10 anos, sem alterar valores."""
    def produzir() -> list[dict]:
        hoje = date.today()
        pontos: list[dict] = []
        ano = desde or INICIO_PADRAO.year
        while ano <= hoje.year:
            fim_ano = min(ano + 9, hoje.year)
            dados = _json_com_nova_tentativa(f"{SGS}.{codigo}/dados", {
                "formato": "json", "dataInicial": f"01/01/{ano}",
                "dataFinal": f"{hoje.day:02d}/{hoje.month:02d}/{hoje.year}" if fim_ano == hoje.year else f"31/12/{fim_ano}"})
            if isinstance(dados, list):
                for p in dados:
                    v = _num(p.get("valor"))
                    if v is not None:  # lacuna da fonte: simplesmente não há ponto
                        pontos.append({"data": _iso(p["data"]), "valor": v})
            ano = fim_ano + 1
        unicos = {p["data"]: p for p in pontos}
        return [unicos[k] for k in sorted(unicos)]
    return _memo(("bcb", codigo, desde), produzir)


def _sidra_desemprego() -> list[dict]:
    def produzir() -> list[dict]:
        linhas = _http.json(f"{SIDRA}/t/6381/n1/all/v/4099/p/all", {"formato": "json"})
        pontos = []
        for r in linhas[1:]:
            v = _num(r.get("V"))
            cod = str(r.get("D3C", ""))
            if v is None or len(cod) != 6:
                continue
            pontos.append({"data": f"{cod[:4]}-{cod[4:]}-01", "valor": v, "referencia": r.get("D3N")})
        return sorted(pontos, key=lambda p: p["data"])
    return _memo(("sidra", 6381), produzir)


def _so_alteracoes(pontos: list[dict]) -> list[dict]:
    """Selic: mantém só os dias em que a meta mudou (+ o último dia observado)."""
    saida: list[dict] = []
    for p in pontos:
        if not saida or p["valor"] != saida[-1]["valor"]:
            saida.append(dict(p))
    if pontos and saida[-1]["data"] != pontos[-1]["data"]:
        saida.append({**pontos[-1], "mantida": True})  # marca “meta mantida até aqui”
    return saida


def serie(chave: str, inicio: str | None = None, fim: str | None = None) -> dict:
    meta = CATALOGO.get(chave)
    if not meta:
        raise KeyError(chave)
    if meta["fonte"] == "sidra":
        pontos = _sidra_desemprego()
    else:  # séries diárias longas: só os últimos anos (o histórico completo não é necessário)
        pontos = _bcb(meta["codigo"], date.today().year - 2 if meta.get("frequencia") == "diária" and not meta.get("alteracoes") else None)
    if meta.get("alteracoes"):
        pontos = _so_alteracoes(pontos)
    if inicio:
        pontos = [p for p in pontos if p["data"] >= inicio]
    if fim:
        pontos = [p for p in pontos if p["data"] <= fim]
    if meta["fonte"] == "bcb":
        url = f"{SGS}.{meta['codigo']}/dados?formato=json"
        fonte = f"Banco Central do Brasil — SGS, série {meta['codigo']}"
    else:
        url = f"{SIDRA}/t/6381/n1/all/v/4099/p/all"
        fonte = "IBGE — SIDRA, tabela 6381 (PNAD Contínua), variável 4099"
    return {
        "id": chave, **{k: meta[k] for k in ("nome", "unidade", "frequencia", "grupo", "tipo", "mede", "interpreta")},
        "nota": meta.get("nota"), "codigo": meta["codigo"], "fonte": fonte, "url_fonte": url,
        "consultado_em": _agora(),
        "periodo_referencia": {"inicio": pontos[0]["data"] if pontos else None, "fim": pontos[-1]["data"] if pontos else None},
        "pontos": pontos,
    }


def _ultimos(chave: str, n: int) -> list[dict]:
    return serie(chave)["pontos"][-n:]


def _soma_12_meses(chave: str) -> tuple[float | None, str | None]:
    """Soma dos 12 últimos fluxos mensais publicados; None se faltar algum mês."""
    pts = _ultimos(chave, 12)
    if len(pts) < 12:
        return None, None
    meses = [int(p["data"][:4]) * 12 + int(p["data"][5:7]) for p in pts]
    if any(b - a != 1 for a, b in zip(meses, meses[1:])):
        return None, None
    return round(sum(p["valor"] for p in pts), 2), pts[-1]["data"]


def _um_ano_antes(chave: str) -> dict | None:
    pts = serie(chave)["pontos"]
    if len(pts) < 13:
        return None
    ult = pts[-1]
    alvo = f"{int(ult['data'][:4]) - 1}{ult['data'][4:]}"
    return next((p for p in pts if p["data"] == alvo), None)


def _situacao(valor: float) -> str:
    return "deficit" if valor > 0 else "superavit"


def situacao_fiscal_brasil() -> dict:
    """Situação fiscal em 4 dimensões para o setor público consolidado (séries do Banco Central)."""
    prim12, fim = _soma_12_meses("primario")
    nom12, _ = _soma_12_meses("nominal")
    jur12, _ = _soma_12_meses("juros")
    prim_pib = _ultimos("primario12", 1)[0]["valor"]
    nom_pib = _ultimos("nominal12", 1)[0]["valor"]
    jur_pib = _ultimos("juros12", 1)[0]["valor"]
    dlsp = _ultimos("dlsp_pib", 1)[0]
    dlsp_rs = _ultimos("dlsp_rs", 1)[0]
    dbgg = _ultimos("dbgg_pib", 1)[0]
    dbgg_ant = _um_ano_antes("dbgg_pib")
    dlsp_ant = _um_ano_antes("dlsp_pib")
    selic = _ultimos("selic", 1)[0]

    itens_fluxo = []
    if prim12 is not None:
        itens_fluxo.append({"id": "primario", "nome": "Primário (sem juros)", "valor_rs": prim12 * 1e6, "pct_pib": prim_pib, "situacao": _situacao(prim12)})
    if nom12 is not None:
        itens_fluxo.append({"id": "nominal", "nome": "Nominal (com juros da dívida)", "valor_rs": nom12 * 1e6, "pct_pib": nom_pib, "situacao": _situacao(nom12)})
    leitura = None
    if prim12 is not None and nom12 is not None:
        if prim12 <= 0 < nom12:
            leitura = "Sem contar os juros, o setor público arrecadou mais do que gastou, mas a conta de juros da dívida tornou o saldo final negativo."
        elif prim12 > 0:
            leitura = "Mesmo sem contar os juros, o setor público gastou mais do que arrecadou, e os juros da dívida aumentam ainda mais o buraco."
        else:
            leitura = "O setor público arrecadou mais do que gastou, mesmo depois de pagar os juros da dívida."

    variacao = round(dbgg["valor"] - dbgg_ant["valor"], 2) if dbgg_ant else None
    tendencia = None if variacao is None else ("alta" if variacao > 0.05 else "queda" if variacao < -0.05 else "estavel")

    custo_impl = round(jur_pib / dlsp["valor"] * 100, 1) if dlsp["valor"] else None
    pib = None
    if dlsp["valor"] and dlsp_rs["data"] == dlsp["data"]:
        pib = {"valor_rs": dlsp_rs["valor"] * 1e6 / (dlsp["valor"] / 100), "data": dlsp["data"],
               "nota": "Derivado do saldo da dívida líquida (R$ ÷ % do PIB) divulgado pelo Banco Central; é o PIB acumulado em 12 meses."}
    try:
        diagnostico = diagnostico_economico()
    except Exception as exc:  # noqa: BLE001 — o radar nunca derruba a situação fiscal
        log.warning("diagnóstico econômico indisponível: %s", exc)
        diagnostico = None
    return {
        "diagnostico": diagnostico,
        "pib": pib,
        "universo": "Setor público consolidado (União, estados, municípios e estatais; Banco Central)",
        "referencia": fim,
        "fluxo": {"periodo": f"12 meses até {fim[5:7]}/{fim[:4]}" if fim else None, "itens": itens_fluxo, "leitura": leitura},
        "estoque": {
            "itens": [
                {"id": "dbgg", "nome": "Dívida bruta do governo geral", "pct": dbgg["valor"], "rotulo_pct": "% do PIB", "data": dbgg["data"],
                 "variacao_pp": variacao},
                {"id": "dlsp", "nome": "Dívida líquida do setor público", "pct": dlsp["valor"], "rotulo_pct": "% do PIB", "data": dlsp["data"],
                 "valor_rs": dlsp_rs["valor"] * 1e6,
                 "variacao_pp": round(dlsp["valor"] - dlsp_ant["valor"], 2) if dlsp_ant else None},
            ],
            "tendencia": tendencia,
            "variacao_pp": variacao,
        },
        "custo": {
            "itens": [
                {"id": "juros12", "nome": "Juros nominais acumulados (12 meses)", "valor_rs": jur12 * 1e6 if jur12 is not None else None,
                 "pct": jur_pib, "rotulo_pct": "% do PIB"},
                {"id": "selic", "nome": "Meta Selic hoje", "pct": selic["valor"], "rotulo_pct": "% ao ano"},
                {"id": "implicito", "nome": "Custo implícito da dívida líquida (aprox.)", "pct": custo_impl, "rotulo_pct": "% ao ano",
                 "nota": "Cálculo nosso: juros dos últimos 12 meses ÷ dívida líquida, ambos em % do PIB."},
            ],
            "leitura": "Quanto da riqueza do país vai apenas para remunerar a dívida. Os juros acompanham a Selic: quando ela sobe, a dívida fica mais cara.",
        },
        "fonte": {"nome": "Banco Central do Brasil — SGS (séries 4649, 4583, 4616, 5793, 5727, 5760, 13762, 4513, 4478, 432)",
                  "url": "https://dadosabertos.bcb.gov.br/"},
    }


# --------------------------------------------------------------------------- #
# Radar de saúde econômica — 4 estágios
# --------------------------------------------------------------------------- #
ESTAGIOS = [
    {"id": "desenvolvimento", "rotulo": "Desenvolvimento", "resumo": "Economia crescendo e contas sob controle."},
    {"id": "recessao", "rotulo": "Recessão", "resumo": "A atividade econômica encolhe por dois trimestres seguidos."},
    {"id": "aperto", "rotulo": "Crise / Aperto", "resumo": "Dívida alta, juros pesados e crédito caro apertam o Estado e as famílias."},
    {"id": "insolvencia", "rotulo": "Risco de insolvência", "resumo": "A dívida cresce sozinha e o país corre risco de não conseguir pagá-la."},
]
# Limites do radar: critério EDITORIAL do projeto (não é métrica oficial); os números vêm das séries oficiais.
# Cada sinal LIGA ao atingir o limite e só DESLIGA depois de recuar a “desliga”: evita que o radar oscile
# quando o indicador anda em cima da linha.
SINAIS_APERTO = {
    "divida": dict(nome="Dívida bruta do governo geral", unidade="% do PIB", liga=75.0, desliga=70.0),
    "nominal": dict(nome="Déficit nominal (12 meses)", unidade="% do PIB", liga=5.0, desliga=4.0),
    "juro_real": dict(nome="Juro real (Selic − IPCA 12 meses)", unidade="p.p.", liga=6.0, desliga=5.0),
}
MIN_SINAIS_APERTO = 2
QUEDA_MIN_RECESSAO = 0.5    # queda acumulada do IBC-Br nos dois trimestres, em %
LIM_DIVIDA_INSOLVENCIA = 100.0


def _fmt(v: float, dec: int = 1) -> str:
    return f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_rs(valor_rs: float) -> str:
    return f"R$ {_fmt(valor_rs / 1e12, 2)} tri" if abs(valor_rs) >= 1e12 else f"R$ {_fmt(valor_rs / 1e9, 0)} bi"


def _mes_antes(mes: str, n: int) -> str:
    idx = int(mes[:4]) * 12 + int(mes[5:7]) - 1 - n
    return f"{idx // 12}-{idx % 12 + 1:02d}"


def _atividade_em(ibc: dict[str, float], mes: str) -> dict | None:
    """IBC-Br: média dos 3 últimos meses contra os 3 anteriores, duas vezes seguidas (trimestres que terminam em `mes`)."""
    medias = []
    for k in range(3):
        v = [ibc.get(_mes_antes(mes, 3 * k + j)) for j in range(3)]
        if None in v:
            return None
        medias.append(sum(v) / 3)
    var = [(medias[0] / medias[1] - 1) * 100, (medias[1] / medias[2] - 1) * 100]
    acum = (medias[0] / medias[2] - 1) * 100
    return {"variacoes_pct": [round(var[1], 2), round(var[0], 2)], "acumulada_pct": round(acum, 2),
            "recessao": var[0] < 0 and var[1] < 0 and acum <= -QUEDA_MIN_RECESSAO}


def _selic_em(pts: list[dict], mes: str) -> float | None:
    v = None
    for p in pts:
        if p["data"][:7] > mes:
            break
        v = p["valor"]
    return v


def _dados_mensais() -> tuple[dict[str, dict[str, float]], list[dict]]:
    """Séries do radar indexadas por mês (AAAA-MM) e a Selic (só as mudanças de meta)."""
    chaves = ("dbgg_pib", "nominal12", "primario12", "ipca12")
    dados = {k: {p["data"][:7]: p["valor"] for p in serie(k)["pontos"]} for k in chaves}
    try:
        dados["ibcbr"] = {p["data"][:7]: p["valor"] for p in serie("ibcbr")["pontos"]}
    except Exception:  # noqa: BLE001 — sem a série o radar segue só com os sinais fiscais
        dados["ibcbr"] = {}
    return dados, serie("selic")["pontos"]


def classificar_historico(desde: str = "2011-01") -> list[dict]:
    """Estágio do radar mês a mês (com a memória dos sinais). O último item é o diagnóstico de hoje;
    o restante serve para conferir o critério contra o que de fato aconteceu."""
    dados, selic = _dados_mensais()
    meses = sorted(m for m in dados["nominal12"]
                   if m >= desde and m in dados["dbgg_pib"] and m in dados["ipca12"] and m in dados["primario12"])
    ligados: set[str] = set()
    saida: list[dict] = []
    for m in meses:
        s_selic = _selic_em(selic, m)
        if s_selic is None:
            continue
        valores = {"divida": dados["dbgg_pib"][m], "nominal": dados["nominal12"][m], "juro_real": round(s_selic - dados["ipca12"][m], 2)}
        sinais = []
        for sid, cfg in SINAIS_APERTO.items():
            v = valores[sid]
            on = v >= cfg["liga"] or (sid in ligados and v >= cfg["desliga"])
            (ligados.add if on else ligados.discard)(sid)
            sinais.append({"id": sid, "nome": cfg["nome"], "valor": v, "unidade": cfg["unidade"], "disparado": on,
                           "regra": f"liga em ≥ {_fmt(cfg['liga'], 0)}, desliga abaixo de {_fmt(cfg['desliga'], 0)}"})
        ativ = _atividade_em(dados["ibcbr"], m)
        antes = dados["dbgg_pib"].get(_mes_antes(m, 12))
        n = sum(x["disparado"] for x in sinais)
        if valores["divida"] >= LIM_DIVIDA_INSOLVENCIA and dados["primario12"][m] > 0 and antes is not None and valores["divida"] > antes:  # BCB: positivo = déficit
            est = "insolvencia"
        elif n >= MIN_SINAIS_APERTO:
            est = "aperto"
        elif ativ and ativ["recessao"]:
            est = "recessao"
        else:
            est = "desenvolvimento"
        saida.append({"mes": m, "estagio": est, "sinais": sinais, "atividade": ativ, "selic": s_selic,
                      "ipca12": dados["ipca12"][m]})
    return saida


def _cotidiano(est: str, v: dict) -> dict:
    """Tradução para a vida real: o que as métricas significam para crédito, serviços públicos e renda."""
    juros_rs = f" (cerca de {_fmt_rs(v['juros_rs'])} em 12 meses)" if v.get("juros_rs") else ""
    credito = {
        "desenvolvimento": f"Com a Selic em {_fmt(v['selic'], 2)}% ao ano e a inflação em {_fmt(v['ipca12'])}%, o crédito tende a ser mais acessível: financiar a casa, o carro ou o negócio pesa menos no orçamento.",
        "recessao": "Com a atividade em queda, bancos e empresas tendem a ficar mais cautelosos: o crédito encurta, a inadimplência costuma subir e investir em um negócio novo fica mais arriscado.",
        "aperto": f"A Selic está em {_fmt(v['selic'], 2)}% ao ano e a inflação em {_fmt(v['ipca12'])}%: sobram {_fmt(v['juro_real'])} pontos de juro real. Na prática, cartão, cheque especial e financiamentos tendem a custar caro, e empresas costumam adiar investimentos e contratações.",
        "insolvencia": "Quando o país perde a confiança de quem lhe empresta, os juros tendem a disparar para todos: o crédito fica escasso ou proibitivo para famílias e empresas.",
    }[est]
    servicos = {
        "desenvolvimento": f"Os juros da dívida consomem {_fmt(v['juros_pib'])}% do PIB{juros_rs}: sobra mais espaço no orçamento para saúde, educação e obras.",
        "recessao": "A arrecadação tende a cair junto com a renda e o emprego, e o governo fica com menos dinheiro para saúde, educação e obras justamente quando mais gente precisa.",
        "aperto": f"Só de juros, o setor público paga {_fmt(v['juros_pib'])}% do PIB{juros_rs}. Cada real que vai para os credores é um real que não está disponível para hospitais, escolas, creches e obras, o que tende a reduzir o espaço para investimentos.",
        "insolvencia": "O orçamento passa a ser engolido pela dívida: sobra pouco para saúde e educação, e cortes ou aumento de impostos tendem a deixar de ser escolha.",
    }[est]
    emprego = f" A taxa de desemprego está em {_fmt(v['desemprego'])}%." if v.get("desemprego") is not None else ""
    renda = {
        "desenvolvimento": "A economia cresce e o ambiente é favorável a vagas e a aumento de renda." + emprego,
        "recessao": "O cenário tende a trazer menos vagas, renda estagnada ou em queda e mais famílias endividadas." + emprego,
        "aperto": f"Com o governo gastando mais do que arrecada, a dívida cresce, e os juros altos mantidos para segurar a inflação ({_fmt(v['ipca12'])}% em 12 meses) tendem a esfriar a economia: menos vagas e menos renda.{emprego}",
        "insolvencia": "O cenário tende a trazer inflação alta, moeda desvalorizada e perda do poder de compra, com risco de um ajuste fiscal abrupto." + emprego,
    }[est]
    titulo = {
        "desenvolvimento": "Na prática: as contas públicas não estão pesando sobre o seu dia a dia",
        "recessao": "Na prática: a economia encolhe, e isso tende a chegar ao emprego e ao crédito",
        "aperto": "Na prática: o país gasta muito para pagar a dívida, e parte da conta tende a chegar ao seu bolso",
        "insolvencia": "Na prática: o país corre o risco de não ter margem para pagar a própria dívida",
    }[est]
    return {"titulo": titulo, "itens": [
        {"tema": "Crédito e juros", "texto": credito},
        {"tema": "Saúde, educação e investimentos", "texto": servicos},
        {"tema": "Emprego e renda", "texto": renda}]}


def diagnostico_economico() -> dict:
    """Em qual dos 4 estágios o país está, com os sinais que levam à conclusão (regras abertas, limites do projeto)."""
    hist = classificar_historico()
    if not hist:
        raise ErroFonte("Sem dados suficientes para o radar de saúde econômica.")
    hoje = hist[-1]
    try:
        desemp = _ultimos("desemprego", 1)[0]["valor"]
    except Exception:  # noqa: BLE001 — IBGE fora do ar: o texto do cotidiano omite a taxa
        desemp = None
    jur12_rs, _ = _soma_12_meses("juros")
    juros_pib = _ultimos("juros12", 1)[0]["valor"]
    # Juros e inflação saem mais cedo que o resultado fiscal: o texto e o sinal usam o dado mais recente de cada um.
    selic = _ultimos("selic", 1)[0]["valor"]
    ipca12 = _ultimos("ipca12", 1)[0]["valor"]
    juro_real = round(selic - ipca12, 2)
    sinais = [{**s, "valor": juro_real} if s["id"] == "juro_real" else s for s in hoje["sinais"]]
    ativ = hoje["atividade"]
    if ativ:
        sinais.append({"id": "atividade", "nome": "Atividade (IBC-Br, últimos 2 trimestres)", "valor": ativ["acumulada_pct"],
                       "unidade": "% acumulado", "disparado": ativ["recessao"],
                       "regra": f"queda em 2 trimestres seguidos, somando ≥ {_fmt(QUEDA_MIN_RECESSAO)}%"})
    idx = [e["id"] for e in ESTAGIOS].index(hoje["estagio"])
    contexto = {"selic": selic, "ipca12": ipca12, "juro_real": juro_real, "juros_pib": juros_pib,
                "juros_rs": jur12_rs * 1e6 if jur12_rs is not None else None, "desemprego": desemp}
    return {
        "estagios": ESTAGIOS, "atual": hoje["estagio"], "indice": idx, "rotulo": ESTAGIOS[idx]["rotulo"],
        "sinais_aperto": sum(s["disparado"] for s in hoje["sinais"]), "sinais": sinais, "atividade": ativ,
        "referencia": f"{hoje['mes']}-01",
        "cotidiano": _cotidiano(hoje["estagio"], contexto),
        "criterio": ("Regra do ImpactoCidadão, não é uma classificação oficial. Crise/Aperto = pelo menos 2 de 3 sinais ligados: dívida bruta "
                     "(liga em 75% do PIB, desliga abaixo de 70%), déficit nominal em 12 meses (liga em 5%, desliga abaixo de 4%) e juro real "
                     "(liga em 6 p.p., desliga abaixo de 5). A margem de saída evita que o radar oscile quando o número anda em cima da linha. "
                     "Recessão = IBC-Br em queda por 2 trimestres seguidos, somando ao menos 0,5%. "
                     "Risco de insolvência = dívida ≥ 100% do PIB, déficit primário e dívida subindo em 12 meses. Vale o estágio mais grave que se aplica."),
        "fonte": {"nome": "Banco Central do Brasil — SGS (séries 13762, 5727, 5760, 5793, 432, 13522, 24364) e IBGE — PNAD Contínua",
                  "url": "https://dadosabertos.bcb.gov.br/"},
    }


def pib_12_meses(ano: int) -> float | None:
    """PIB acumulado em 12 meses (R$) até dezembro de `ano`, derivado do saldo da dívida líquida em R$ e em % do PIB."""
    rs = {p["data"]: p["valor"] for p in serie("dlsp_rs")["pontos"]}
    pc = {p["data"]: p["valor"] for p in serie("dlsp_pib")["pontos"]}
    chave = f"{ano}-12-01"
    if chave in rs and chave in pc and pc[chave]:
        return rs[chave] * 1e6 / (pc[chave] / 100)
    return None


def _soma(valores: list[float | None]) -> float | None:
    return None if any(v is None for v in valores) else round(sum(valores), 2)


def _comparar(serie_mes: dict[str, float | None], fluxo: bool) -> dict | None:
    """Mês × mês anterior e, para fluxos, últimos 12 meses × 12 meses anteriores; só com dados completos."""
    meses = sorted(m for m, v in serie_mes.items() if v is not None)
    if not meses:
        return None
    ult = meses[-1]
    ano, mes = int(ult[:4]), int(ult[5:7])

    def mes_antes(n: int) -> str:
        idx = ano * 12 + mes - 1 - n
        return f"{idx // 12}-{idx % 12 + 1:02d}"

    atual, anterior = serie_mes[ult], serie_mes.get(mes_antes(1))
    out: dict[str, Any] = {"mes": ult, "valor": atual, "mes_anterior": anterior,
                           "mesmo_mes_ano_anterior": serie_mes.get(mes_antes(12))}
    if fluxo:
        a = _soma([serie_mes.get(mes_antes(i)) for i in range(0, 12)])
        b = _soma([serie_mes.get(mes_antes(i)) for i in range(12, 24)])
        out.update({"acumulado_12m": a, "acumulado_12m_anterior": b})
    return out


def financiamento() -> dict:
    """Necessidade de financiamento (resultado fiscal) × emissões, resgates e estoque da DPF."""
    consultado = _agora()
    erros: list[str] = []
    fiscal: dict[str, dict[str, float]] = {}
    for chave in ("primario", "nominal", "juros"):
        try:
            fiscal[chave] = {p["data"][:7]: p["valor"] for p in serie(chave)["pontos"]}
        except Exception as exc:  # noqa: BLE001
            log.warning("fiscal %s indisponível: %s", chave, exc)
            fiscal[chave] = {}
            erros.append("Séries fiscais do Banco Central indisponíveis no momento.")
            break
    try:
        er = _emissoes_resgates()
        emi = {x["mes"]: x for x in er["serie"]}
    except Exception as exc:  # noqa: BLE001
        log.warning("emissões indisponíveis: %s", exc)
        er, emi = None, {}
        erros.append("Emissões e resgates do Tesouro indisponíveis no momento.")
    try:
        dpf = _estoque_dpf()
        extra = dpf["_extra"]
        est = {x["mes"]: x for x in extra["por_mes"]}
    except Exception as exc:  # noqa: BLE001
        log.warning("estoque DPF indisponível: %s", exc)
        dpf, extra, est = None, None, {}
        erros.append("Estoque da Dívida Pública Federal indisponível no momento.")

    todos = sorted(set(fiscal.get("nominal", {})) | set(emi) | set(est))
    todos = todos[-48:]
    meses = [m for m in todos]
    linhas = []
    estoque_mes: dict[str, float | None] = {}
    for m in meses:
        e = emi.get(m) or {}
        em, rs = e.get("emissoes"), e.get("resgates")
        total = est[m]["total"] / 1e6 if m in est else None  # R$ → R$ milhões
        estoque_mes[m] = total
        linhas.append({"mes": m,
                       "primario": fiscal.get("primario", {}).get(m), "nominal": fiscal.get("nominal", {}).get(m),
                       "juros": fiscal.get("juros", {}).get(m),
                       "emissoes": em, "resgates": rs,
                       "emissao_liquida": round(em - rs, 2) if em is not None and rs is not None else None,
                       "estoque": round(total, 2) if total is not None else None})
    for i, l in enumerate(linhas):
        ant = linhas[i - 1]["estoque"] if i else None
        l["variacao_estoque"] = round(l["estoque"] - ant, 2) if l["estoque"] is not None and ant is not None else None
        # diferença entre a variação do estoque e a emissão líquida: juros apropriados, câmbio e outros ajustes
        l["diferenca_estoque_vs_emissao"] = (round(l["variacao_estoque"] - l["emissao_liquida"], 2)
                                             if l["variacao_estoque"] is not None and l["emissao_liquida"] is not None else None)

    def col(chave: str) -> dict[str, float | None]:
        return {l["mes"]: l[chave] for l in linhas}

    comparativos = {
        "primario": _comparar(col("primario"), True), "nominal": _comparar(col("nominal"), True),
        "juros": _comparar(col("juros"), True), "emissoes": _comparar(col("emissoes"), True),
        "resgates": _comparar(col("resgates"), True), "emissao_liquida": _comparar(col("emissao_liquida"), True),
        "estoque": _comparar(col("estoque"), False), "variacao_estoque": _comparar(col("variacao_estoque"), True),
    }

    saida: dict[str, Any] = {
        "consultado_em": consultado, "erros": erros, "serie": linhas, "comparativos": comparativos,
        "nfsp": {
            "indicador": "Resultado nominal do setor público (NFSP, sem desvalorização cambial)",
            "serie_bcb": 4583, "metodologia": "“Abaixo da linha”: variação nominal do saldo da dívida líquida, deduzidos ajustes patrimoniais e metodológicos. "
                                              "O primário é medido por caixa e os juros nominais por competência.",
            "abrangencia": "Governo Federal (Tesouro e Previdência), governos estaduais e municipais, empresas estatais das três esferas (exceto Petrobras e Eletrobras) e Banco Central.",
            "sinal": "No Banco Central, valor positivo = necessidade de financiamento (déficit); negativo = superávit.",
            "url": "https://www3.bcb.gov.br/sgspub/consultarmetadados/consultarMetadadosSeries.do?method=consultarMetadadosSeriesInternet&hdOidSerieSelecionada=4583",
        },
        "fontes": {
            "fiscal": {"nome": "Banco Central do Brasil — SGS (4649, 4583, 4616)", "periodicidade": "mensal", "unidade": "R$ milhões (fluxo mensal corrente)"},
            "emissoes": ({"nome": "Tesouro Transparente — Emissões e Resgates da DPF", "periodicidade": "mensal", "unidade": "R$ milhões",
                          "atualizado": er["atualizado"], "url": er["pagina"]} if er else None),
            "estoque": ({"nome": "Tesouro Transparente — Estoque da DPF", "periodicidade": "mensal", "unidade": "R$ (convertido para R$ milhões nos gráficos)",
                         "atualizado": dpf["atualizado"], "url": dpf["pagina"]} if dpf else None),
        },
    }
    if dpf and extra:
        ref = dpf["referencia"]
        ult = est[ref]
        faixas, tot = extra["faixas"], dpf["total"]
        def faixa(k: str) -> dict:
            v = faixas.get(k, 0.0)
            return {"valor": v, "pct": round(v / tot * 100, 1) if tot else None}
        prazo = extra["prazo_medio_anos"]
        ano_ant = f"{int(ref[:4]) - 1}{ref[4:]}"
        saida["estoque"] = {
            "referencia": ref, "total": tot, "dpmfi": ult["interna"], "dpfe": ult["externa"],
            "pct_dpmfi": round(ult["interna"] / tot * 100, 1), "pct_dpfe": round(ult["externa"] / tot * 100, 1),
            "mercado": ult["mercado"], "banco_central": ult["banco_central"],
        }
        saida["vencimentos"] = {
            "referencia": ref, "curto": faixa("curto"), "medio": faixa("medio"), "longo": faixa("longo"), "sem_data": faixa("sem data"),
            "proximos_12_meses": extra["calendario"],
            "nota": "Valores de face do estoque por data de vencimento; não incluem emissões futuras. Observado, não é projeção.",
        }
        saida["composicao"] = {
            "referencia": ref,
            "indexadores": [{**i, "pct": round(i["valor"] / tot * 100, 1)} for i in extra["indexadores"]],
            "prazo_medio_anos": prazo.get(ref), "prazo_medio_anos_ano_anterior": prazo.get(ano_ant),
            "nota": "Indexador classificado por tipo de título (cálculo nosso) e prazo médio ponderado pelo estoque (cálculo nosso). "
                    "Os números oficiais estão no Relatório Mensal da Dívida do Tesouro.",
        }
    saida["custo_medio"] = {
        "disponivel": False,
        "motivo": "O custo médio da DPF é divulgado no Relatório Mensal da Dívida (PDF) e não há série aberta por API; não estimamos.",
        "url": "https://www.tesourotransparente.gov.br/publicacoes/relatorio-mensal-da-divida-rmd",
    }
    return saida


def reservas() -> dict:
    """Reservas internacionais (saldo) e transações do mês de ouro, DES e posição no FMI."""
    caixa = serie("reservas_caixa")["pontos"]
    liquidez = serie("reservas_liquidez", inicio=(date.today().replace(day=1)).replace(year=date.today().year - 1).isoformat())["pontos"]
    fluxos = []
    for chave in ("fmi_posicao_fluxo", "des_fluxo", "ouro_fluxo"):
        s = serie(chave)
        if s["pontos"]:
            ult = s["pontos"][-1]
            fluxos.append({"id": chave, "nome": s["nome"], "valor": ult["valor"], "data": ult["data"], "unidade": s["unidade"]})
    return {
        "consultado_em": _agora(),
        "caixa": {"valor": caixa[-1]["valor"], "data": caixa[-1]["data"],
                  "anterior": caixa[-2]["valor"] if len(caixa) > 1 else None, "serie": caixa[-36:],
                  "fonte": "Banco Central do Brasil — SGS, série 3546",
                  "url_fonte": f"{SGS}.3546/dados?formato=json"},
        "liquidez": ({"valor": liquidez[-1]["valor"], "data": liquidez[-1]["data"],
                      "url_fonte": f"{SGS}.13982/dados?formato=json"} if liquidez else None),
        "fluxos": fluxos, "nota_fluxos": NOTA_FLUXO_BPM,
        "links": {"fmi": "https://www.imf.org/external/np/fin/tad/exfin2.aspx?memberKey1=90",
                  "bcb": "https://www.bcb.gov.br/estabilidadefinanceira/reservasinternacionais"},
    }


def catalogo() -> list[dict]:
    return [{"id": k, **{c: v[c] for c in ("nome", "unidade", "frequencia", "grupo", "tipo", "mede", "interpreta")},
             "codigo": v["codigo"], "nota": v.get("nota")} for k, v in CATALOGO.items()]


def resumo() -> dict:
    """Valor mais recente de cada indicador (com o anterior, para a variação)."""
    cartoes, erros = [], []
    for chave in CATALOGO:
        try:
            s = serie(chave)
        except Exception as exc:  # noqa: BLE001 — uma fonte fora do ar não derruba as outras
            log.warning("série %s indisponível: %s", chave, exc)
            erros.append({"id": chave, "erro": "Fonte indisponível no momento."})
            continue
        if not s["pontos"]:
            continue
        ult = s["pontos"][-1]
        ant = s["pontos"][-2] if len(s["pontos"]) > 1 else None
        cartoes.append({"id": chave, "nome": s["nome"], "unidade": s["unidade"], "grupo": s["grupo"],
                        "frequencia": s["frequencia"], "valor": ult["valor"], "data": ult["data"],
                        "referencia": ult.get("referencia"), "anterior": ant["valor"] if ant else None,
                        "fonte": s["fonte"], "url_fonte": s["url_fonte"]})
    return {"consultado_em": _agora(), "cartoes": cartoes, "erros": erros}


# --------------------------------------------------------------------------- #
# Tesouro Transparente (CKAN → CSV/XLSX)
# --------------------------------------------------------------------------- #
def _recurso(pacote: str, formato: str) -> dict:
    p = _http.json(f"{CKAN}/package_show", {"id": pacote})
    recursos = [r for r in p["result"]["resources"] if (r.get("format") or "").upper() == formato]
    if not recursos:
        raise ErroFonte(f"Conjunto {pacote} sem arquivo {formato}")
    return {"url": recursos[0]["url"], "atualizado": p["result"].get("metadata_modified"),
            "titulo": p["result"].get("title"), "pagina": PAGINA_TESOURO + pacote}


def _csv(pacote: str) -> tuple[list[list[str]], dict]:
    rec = _recurso(pacote, "CSV")
    bruto = _http_arquivos.get(rec["url"], accept="*/*").content
    try:
        texto = bruto.decode("utf-8")
    except UnicodeDecodeError:
        texto = bruto.decode("latin-1")
    leitor = csv.reader(io.StringIO(texto), delimiter=";")
    next(leitor)
    return list(leitor), rec


def _br(txt: str) -> float | None:
    try:
        return float(txt.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def _mes_iso(mm_aaaa: str) -> str:
    mm, aaaa = mm_aaaa.split("/")
    return f"{aaaa}-{mm}"


def _grupo_titulo(nome: str) -> str:
    m = re.match(r"^[A-Za-zÀ-ú]+(?:-[A-Za-z]+)?", nome.strip())
    return m.group(0).upper() if m else "OUTROS"


# Indexador por tipo de título (classificação nossa; o Relatório Mensal da Dívida do Tesouro é a referência oficial)
_INDEXADOR = {
    "LFT": "Selic (pós-fixado)", "LTN": "Prefixado", "NTN-F": "Prefixado",
    "NTN-B": "Índice de preços (IPCA/IGP-M)", "NTN-C": "Índice de preços (IPCA/IGP-M)", "NTN-I": "Índice de preços (IPCA/IGP-M)",
    "NTN-D": "Câmbio",
}


def _indexador(titulo: str, tipo_divida: str) -> str:
    if "Externa" in tipo_divida:
        return "Câmbio (dívida externa)"
    return _INDEXADOR.get(_grupo_titulo(titulo), "Demais (TDA, contratos e outros)")


def _estoque_dpf() -> dict:
    def produzir() -> dict:
        linhas, rec = _csv("estoque-da-divida-publica-federal")
        total: dict[str, float] = defaultdict(float)
        por_classe: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        por_tipo: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        por_titulo: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        venc: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        peso_prazo: dict[str, float] = defaultdict(float)       # Σ valor × anos até o vencimento
        base_prazo: dict[str, float] = defaultdict(float)       # Σ valor com vencimento válido
        por_indexador: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        faixas: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        calendario: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))  # mês de referência → AAAA-MM de vencimento
        for r in linhas:
            if len(r) < 7:
                continue
            v = _br(r[2])
            if v is None:
                continue
            mes = _mes_iso(r[4])
            total[mes] += v
            por_classe[mes][r[5]] += v
            por_tipo[mes][r[6]] += v
            por_titulo[mes][_grupo_titulo(r[0])] += v
            por_indexador[mes][_indexador(r[0], r[6])] += v
            m = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", r[1].strip())
            venc[mes][m.group(3) if m else "s/d"] += v
            if not m:
                faixas[mes]["sem data"] += v
                continue
            vm = int(m.group(3)) * 12 + int(m.group(2))
            rm = int(mes[:4]) * 12 + int(mes[5:7])
            meses_ate = vm - rm
            peso_prazo[mes] += v * max(meses_ate, 0) / 12
            base_prazo[mes] += v
            faixas[mes]["curto" if meses_ate <= 12 else "medio" if meses_ate <= 36 else "longo"] += v
            if 0 <= meses_ate <= 12:
                calendario[mes][f"{m.group(3)}-{m.group(2)}"] += v
        meses = sorted(total)
        ultimo = meses[-1]
        titulos = sorted(por_titulo[ultimo].items(), key=lambda kv: -kv[1])
        principais, resto = titulos[:8], sum(v for _, v in titulos[8:])
        composicao = [{"nome": n, "valor": round(v, 2)} for n, v in principais]
        if resto:
            composicao.append({"nome": "Demais", "valor": round(resto, 2)})
        perfil = [{"ano": a, "valor": round(v, 2)} for a, v in sorted(venc[ultimo].items())]
        prazo = {m: round(peso_prazo[m] / base_prazo[m], 2) for m in meses if base_prazo[m]}
        indexadores = [{"nome": n, "valor": round(v, 2)} for n, v in sorted(por_indexador[ultimo].items(), key=lambda kv: -kv[1])]
        return {
            "_extra": {
                "por_mes": [{"mes": m, "total": round(total[m], 2),
                             "interna": round(sum(v for k, v in por_tipo[m].items() if "Interna" in k), 2),
                             "externa": round(sum(v for k, v in por_tipo[m].items() if "Externa" in k), 2),
                             "mercado": round(por_classe[m].get("Mercado", 0), 2),
                             "banco_central": round(por_classe[m].get("Banco Central", 0), 2)} for m in meses],
                "prazo_medio_anos": prazo,
                "faixas": {k: round(v, 2) for k, v in faixas[ultimo].items()},
                "calendario": [{"mes": k, "valor": round(v, 2)} for k, v in sorted(calendario[ultimo].items())],
                "indexadores": indexadores,
            },
            "referencia": ultimo, "atualizado": rec["atualizado"], "pagina": rec["pagina"], "url": rec["url"],
            "evolucao": [{"mes": m, "valor": round(total[m], 2)} for m in meses],
            "total": round(total[ultimo], 2),
            "por_tipo": [{"nome": n, "valor": round(v, 2)} for n, v in por_tipo[ultimo].items()],
            "por_classe": [{"nome": n, "valor": round(v, 2)} for n, v in por_classe[ultimo].items()],
            "composicao": composicao, "vencimentos": perfil,
            "total_por_mes": {m: total[m] for m in meses},
        }
    return _memo(("dpf",), produzir)


def _emissoes_resgates() -> dict:
    def produzir() -> dict:
        from openpyxl import load_workbook

        rec = _recurso("emissoes-e-resgates-divida-publica-federal", "XLSX")
        bruto = _http_arquivos.get(rec["url"], accept="*/*").content
        ws = load_workbook(io.BytesIO(bruto), read_only=True, data_only=True).worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        cab = next(r for r in linhas if sum(bool(isinstance(c, str) and re.fullmatch(r"[A-Za-zç]{3}/\d{2}", c.strip())) for c in r) > 12)
        meses_pt = {"Jan": 1, "Fev": 2, "Mar": 3, "Abr": 4, "Mai": 5, "Jun": 6, "Jul": 7, "Ago": 8, "Set": 9, "Out": 10, "Nov": 11, "Dez": 12}
        colunas: list[tuple[int, str]] = []
        for i, c in enumerate(cab):
            if isinstance(c, str) and re.fullmatch(r"[A-Za-zç]{3}/\d{2}", c.strip()):
                mes, ano = c.strip().split("/")
                if mes.capitalize() in meses_pt:
                    colunas.append((i, f"20{ano}-{meses_pt[mes.capitalize()]:02d}"))

        def linha(rotulo: str, ocorrencia: int = 0) -> dict[str, float | None]:
            achadas = [r for r in linhas if isinstance(r[0], str) and r[0].strip().startswith(rotulo)]
            r = achadas[ocorrencia]
            return {m: (float(r[i]) if isinstance(r[i], (int, float)) else None) for i, m in colunas}

        def linha_exata(rotulo: str) -> dict[str, float | None]:
            try:
                return linha(rotulo)
            except IndexError:
                return {}

        emissoes, resgates = linha("EMISSÕES"), linha("RESGATES")
        em_dpmfi, em_dpfe = linha_exata("I - DPMFi"), linha_exata("II - DPFe")
        rs_dpmfi, rs_dpfe = linha_exata("III - DPMFi"), linha_exata("IV - DPFe")
        liq_dpmfi, liq_dpfe = linha_exata("DPMFi    (I - III)"), linha_exata("DPFe    (II - IV)")
        meses = [m for _, m in colunas]
        return {
            "unidade": "R$ milhões (conforme a planilha do Tesouro)", "atualizado": rec["atualizado"],
            "pagina": rec["pagina"], "url": rec["url"],
            "serie": [{"mes": m, "emissoes": emissoes.get(m), "resgates": resgates.get(m),
                       "emissoes_dpmfi": em_dpmfi.get(m), "emissoes_dpfe": em_dpfe.get(m),
                       "resgates_dpmfi": rs_dpmfi.get(m), "resgates_dpfe": rs_dpfe.get(m),
                       "liquida_dpmfi": liq_dpmfi.get(m), "liquida_dpfe": liq_dpfe.get(m)} for m in meses
                      if emissoes.get(m) is not None or resgates.get(m) is not None],
        }
    return _memo(("emissoes",), produzir)


def _tesouro_direto() -> dict:
    def produzir() -> dict:
        linhas, rec_e = _csv("estoque-do-tesouro-direto")
        total: dict[str, float] = defaultdict(float)
        por_tipo: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for r in linhas:
            if len(r) < 6:
                continue
            v = _br(r[5])
            if v is None:
                continue
            mes = _mes_iso(r[2])
            total[mes] += v
            por_tipo[mes][r[0]] += v
        meses = sorted(total)
        ultimo = meses[-1]
        taxas_linhas, rec_t = _csv("taxas-dos-titulos-ofertados-pelo-tesouro-direto")

        def chave_data(d: str) -> str:
            dia, mes, ano = d.split("/")
            return f"{ano}-{mes}-{dia}"

        data_base = max((r[2] for r in taxas_linhas if len(r) >= 7), key=chave_data)
        taxas = []
        for r in taxas_linhas:
            if len(r) >= 7 and r[2] == data_base:
                taxas.append({"titulo": r[0], "vencimento": _iso(r[1]), "taxa_compra": _br(r[3]), "taxa_venda": _br(r[4]),
                              "pu_compra": _br(r[5]), "pu_venda": _br(r[6])})
        taxas.sort(key=lambda t: (t["titulo"], t["vencimento"]))
        return {
            "estoque": {"referencia": ultimo, "total": round(total[ultimo], 2),
                        "por_tipo": [{"nome": n, "valor": round(v, 2)} for n, v in sorted(por_tipo[ultimo].items(), key=lambda kv: -kv[1])],
                        "evolucao": [{"mes": m, "valor": round(total[m], 2)} for m in meses],
                        "atualizado": rec_e["atualizado"], "pagina": rec_e["pagina"], "url": rec_e["url"],
                        "total_por_mes": {m: total[m] for m in meses}},
            "taxas": {"data_base": _iso(data_base), "titulos": taxas, "atualizado": rec_t["atualizado"],
                      "pagina": rec_t["pagina"], "url": rec_t["url"]},
        }
    return _memo(("td",), produzir)


def divida() -> dict:
    """Dívida Pública Federal (bruta) e, separadamente, a parcela do Tesouro Direto."""
    partes: dict[str, Any] = {"consultado_em": _agora(), "erros": []}

    try:
        d = _estoque_dpf()
        partes["dpf"] = {k: v for k, v in d.items() if k not in ("total_por_mes", "_extra")}
    except Exception as exc:  # noqa: BLE001
        log.warning("DPF indisponível: %s", exc)
        d = None
        partes["erros"].append("Estoque da Dívida Pública Federal indisponível no momento.")

    try:
        e = _emissoes_resgates()
        partes["emissoes_resgates"] = e
    except Exception as exc:  # noqa: BLE001
        log.warning("Emissões/resgates indisponíveis: %s", exc)
        partes["erros"].append("Emissões e resgates da DPF indisponíveis no momento.")

    try:
        td = _tesouro_direto()
        est = td["estoque"]
        partes["tesouro_direto"] = {"estoque": {k: v for k, v in est.items() if k != "total_por_mes"}, "taxas": td["taxas"]}
        if d and est["referencia"] in d["total_por_mes"]:
            ref = est["referencia"]
            partes["tesouro_direto"]["participacao_na_dpf"] = {
                "mes": ref, "tesouro_direto": est["total_por_mes"][ref], "dpf": d["total_por_mes"][ref],
                "pct": round(100 * est["total_por_mes"][ref] / d["total_por_mes"][ref], 2)}
    except Exception as exc:  # noqa: BLE001
        log.warning("Tesouro Direto indisponível: %s", exc)
        partes["erros"].append("Dados do Tesouro Direto indisponíveis no momento.")
    return partes
